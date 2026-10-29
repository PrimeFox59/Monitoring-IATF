from flask import Blueprint, request, jsonify, send_file
from models import db, User
from routes.audit import log_activity
from datetime import datetime
import io
import openpyxl
from openpyxl.styles import Font, PatternFill

users_bp = Blueprint('users', __name__)

@users_bp.route('/api/login', methods=['POST'])
def login():
    data = request.get_json(silent=True) or request.json or {}
    user_id = (data.get('id') or '').strip()
    password = (data.get('password') or '').strip()

    user = User.query.filter_by(id=user_id).first()
    if user:
        # Flexible authentication for demo: match password or master admin pass
        valid_pass = (user.password == password) or (user.id.lower() == 'admin' and password in ('password', 'admin', '327856'))
        if not valid_pass:
            user = None
    if user:
        if not user.is_approved:
            return jsonify({'error': 'Akun Anda belum disetujui oleh Admin'}), 401
        log_activity(user.id, user.full_name, 'LOGIN', 'User logged in')
        return jsonify({
            'user': user.to_dict(),
            'message': 'Login successful'
        }), 200
    return jsonify({'error': 'ID atau Password salah'}), 401

@users_bp.route('/api/register', methods=['POST'])
def register():
    """Self-registration for new users"""
    try:
        data = request.get_json(silent=True) or request.json or {}
        user_id = data.get('id')

        if not user_id:
            return jsonify({'error': 'User ID wajib diisi'}), 400

        if User.query.get(user_id):
            return jsonify({'error': 'User ID sudah terdaftar'}), 400

        new_user = User(
            id=user_id,
            password=data.get('password'),
            full_name=data.get('full_name'),
            department=data.get('department'),
            section=data.get('section'),
            role='Staff',
            is_approved=0
        )
        db.session.add(new_user)
        db.session.commit()

        log_activity(
            new_user.id,
            new_user.full_name,
            'REGISTER',
            f'User registered: {new_user.full_name}'
        )

        return jsonify({
            'message': 'Registrasi berhasil! Mohon tunggu persetujuan Admin.',
            'user': new_user.to_dict()
        }), 201

    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500

@users_bp.route('/api/users', methods=['GET'])
def get_users():
    user_role = request.headers.get('X-User-Role')
    if not user_role or user_role.lower() not in ('admin', 'administrator'):
        return jsonify({'error': 'Unauthorized: Admin role required to list users'}), 403

    users = User.query.all()
    return jsonify([u.to_dict() for u in users]), 200

@users_bp.route('/api/users', methods=['POST'])
def create_user():
    user_role = request.headers.get('X-User-Role')
    if not user_role or user_role.lower() not in ('admin', 'administrator'):
        return jsonify({'error': 'Unauthorized: Admin role required to create users'}), 403

    data = request.get_json(silent=True) or request.json or {}
    if User.query.get(data.get('id')):
        return jsonify({'error': 'User ID sudah terdaftar'}), 400

    new_user = User(
        id=data.get('id'),
        password=data.get('password'),
        full_name=data.get('full_name'),
        department=data.get('department'),
        section=data.get('section'),
        role=data.get('role', 'Staff'),
        is_approved=data.get('is_approved', 1)
    )
    db.session.add(new_user)
    db.session.commit()

    admin_name = request.headers.get('X-User-Name', 'Admin')
    log_activity(
        request.headers.get('X-User-Role', 'System'),
        admin_name,
        'CREATE_USER',
        f'Created user: {new_user.id} ({new_user.full_name})'
    )

    return jsonify(new_user.to_dict()), 201

@users_bp.route('/api/users/<user_id>', methods=['PUT'])
def update_user(user_id):
    user_role = request.headers.get('X-User-Role')
    if not user_role or user_role.lower() not in ('admin', 'administrator'):
        return jsonify({'error': 'Unauthorized: Admin role required to update users'}), 403

    user = User.query.get_or_404(user_id)
    data = request.get_json(silent=True) or request.json or {}

    if 'full_name' in data:
        user.full_name = data['full_name']
    if 'password' in data and data['password']:
        user.password = data['password']
    if 'department' in data:
        user.department = data['department']
    if 'section' in data:
        user.section = data['section']
    if 'role' in data:
        user.role = data['role']
    if 'is_approved' in data:
        user.is_approved = 1 if data['is_approved'] else 0

    db.session.commit()

    admin_name = request.headers.get('X-User-Name', 'Admin')
    log_activity(
        request.headers.get('X-User-Role', 'System'),
        admin_name,
        'UPDATE_USER',
        f'Updated user: {user_id}'
    )

    return jsonify(user.to_dict()), 200

@users_bp.route('/api/users/<user_id>', methods=['DELETE'])
def delete_user(user_id):
    user_role = request.headers.get('X-User-Role')
    if not user_role or user_role.lower() not in ('admin', 'administrator'):
        return jsonify({'error': 'Unauthorized: Admin role required to delete users'}), 403

    user = User.query.get_or_404(user_id)
    user_name_deleted = user.full_name
    db.session.delete(user)
    db.session.commit()

    admin_name = request.headers.get('X-User-Name', 'Admin')
    log_activity(
        request.headers.get('X-User-Role', 'System'),
        admin_name,
        'DELETE_USER',
        f'Deleted user: {user_id} ({user_name_deleted})'
    )

    return jsonify({'message': 'User deleted'}), 200

@users_bp.route('/api/profile', methods=['PUT'])
def update_profile():
    data = request.get_json(silent=True) or request.json or {}
    user_id = data.get('id')
    user = User.query.get_or_404(user_id)

    if 'full_name' in data:
        user.full_name = data['full_name']
    if 'section' in data:
        user.section = data['section']
    if 'password' in data and data['password']:
        user.password = data['password']

    db.session.commit()
    log_activity(user.id, user.full_name, 'UPDATE_PROFILE', 'User updated their profile')

    return jsonify(user.to_dict()), 200

@users_bp.route('/api/users/template', methods=['GET'])
def download_user_template():
    """Download Excel template for bulk user upload"""
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Users'

        headers = ['id', 'password', 'full_name', 'department', 'section', 'role', 'is_approved']
        ws.append(headers)

        example = ['user123', 'password123', 'John Doe', 'Engineering', 'QA', 'Staff', 'TRUE']
        ws.append(example)

        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font

        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 20
        ws.column_dimensions['F'].width = 10
        ws.column_dimensions['G'].width = 12

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='user_upload_template.xlsx'
        )

    except Exception as e:
        return jsonify({'error': str(e)}), 500



@users_bp.route('/api/users/upload', methods=['POST'])
def upload_users():
    user_role = request.headers.get('X-User-Role')
    if not user_role or user_role.lower() not in ('admin', 'administrator'):
        return jsonify({'error': 'Unauthorized: Admin role required to upload users'}), 403

    if 'file' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400

    file = request.files['file']
    try:
        wb = openpyxl.load_workbook(file)
        ws = wb.active
        created = []
        errors = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue
            user_id = str(row[0]).strip() if row[0] is not None else ''
            if not user_id:
                continue

            password = str(row[1]).strip() if len(row) > 1 and row[1] is not None else 'password'
            full_name = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ''
            dept = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ''
            sec = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ''
            role = str(row[5]).strip() if len(row) > 5 and row[5] is not None else 'Staff'
            is_app = 1 if len(row) > 6 and str(row[6]).upper() in ('TRUE', '1', 'YES') else 1

            if User.query.get(user_id):
                errors.append(f"User ID {user_id} already exists")
                continue

            new_user = User(
                id=user_id,
                password=password,
                full_name=full_name,
                department=dept,
                section=sec,
                role=role,
                is_approved=is_app
            )
            db.session.add(new_user)
            created.append(user_id)

        db.session.commit()
        return jsonify({'created': created, 'errors': errors}), 200

    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500
