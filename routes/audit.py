from flask import Blueprint, jsonify, request, send_file
from models import db, AuditLog
from datetime import datetime
import json
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

audit_bp = Blueprint('audit', __name__)

def log_activity(user_id, user_name, action, details):
    """
    Helper function to log user activity.
    details can be a dictionary or string.
    """
    try:
        if isinstance(details, (dict, list)):
            details_str = json.dumps(details)
        else:
            details_str = str(details)
            
        log = AuditLog(
            timestamp=datetime.now().isoformat(),
            user_id=str(user_id),
            user_name=user_name,
            action=action,
            details=details_str
        )
        db.session.add(log)
        db.session.commit()
    except Exception as e:
        print(f"Failed to log activity: {e}")

@audit_bp.route('/api/audit-logs', methods=['GET'])
def get_audit_logs():
    """Get audit logs with optional filtering"""
    user_role = request.headers.get('X-User-Role')
    if not user_role or user_role.lower() not in ('admin', 'administrator'):
        return jsonify({'error': 'Unauthorized: Admin role required to view audit logs'}), 403

    try:
        page = request.args.get('page', 1, type=int)
        per_page = request.args.get('per_page', 50, type=int)
        search = request.args.get('search', '').strip()
        start_date = request.args.get('start_date', '').strip()
        end_date = request.args.get('end_date', '').strip()

        query = AuditLog.query.order_by(AuditLog.timestamp.desc())

        if search:
            search_term = f"%{search}%"
            query = query.filter(
                (AuditLog.user_name.ilike(search_term)) |
                (AuditLog.action.ilike(search_term)) |
                (AuditLog.details.ilike(search_term))
            )

        if start_date:
            try:
                start_datetime = datetime.strptime(start_date, '%Y-%m-%d').replace(hour=0, minute=0, second=0)
                start_iso = start_datetime.isoformat()
                query = query.filter(AuditLog.timestamp >= start_iso)
            except ValueError:
                return jsonify({'error': 'Invalid start_date format. Use YYYY-MM-DD'}), 400

        if end_date:
            try:
                end_datetime = datetime.strptime(end_date, '%Y-%m-%d').replace(hour=23, minute=59, second=59)
                end_iso = end_datetime.isoformat()
                query = query.filter(AuditLog.timestamp <= end_iso)
            except ValueError:
                return jsonify({'error': 'Invalid end_date format. Use YYYY-MM-DD'}), 400

        paginated = query.paginate(page=page, per_page=per_page, error_out=False)

        return jsonify({
            'logs': [log.to_dict() for log in paginated.items],
            'total': paginated.total,
            'pages': paginated.pages,
            'current_page': page
        }), 200

    except Exception as e:
        return jsonify({'error': str(e)}), 500

@audit_bp.route('/api/audit-logs/export', methods=['GET'])
def export_audit_logs():
    """Export filtered audit logs to a professionally styled Excel file"""
    user_role = request.args.get('user_role') or request.headers.get('X-User-Role')
    if not user_role or user_role.lower() not in ('admin', 'administrator'):
        return jsonify({'error': 'Unauthorized: Admin role required to export audit logs'}), 403

    try:
        search = request.args.get('search', '').strip()
        start_date = request.args.get('start_date', '').strip()
        end_date = request.args.get('end_date', '').strip()

        query = AuditLog.query.order_by(AuditLog.timestamp.desc())

        if search:
            search_term = f"%{search}%"
            query = query.filter(
                (AuditLog.user_name.ilike(search_term)) |
                (AuditLog.action.ilike(search_term)) |
                (AuditLog.details.ilike(search_term))
            )

        if start_date:
            try:
                start_datetime = datetime.strptime(start_date, '%Y-%m-%d').replace(hour=0, minute=0, second=0)
                query = query.filter(AuditLog.timestamp >= start_datetime.isoformat())
            except ValueError:
                pass

        if end_date:
            try:
                end_datetime = datetime.strptime(end_date, '%Y-%m-%d').replace(hour=23, minute=59, second=59)
                query = query.filter(AuditLog.timestamp <= end_datetime.isoformat())
            except ValueError:
                pass

        logs = query.all()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Audit Trail Proyek"
        ws.views.sheetView[0].showGridLines = True

        navy_fill = PatternFill(start_color="18345E", end_color="18345E", fill_type="solid")
        accent_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        title_font = Font(name="Segoe UI", size=14, bold=True, color="18345E")
        subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        data_font = Font(name="Segoe UI", size=10, color="1E293B")
        bold_data_font = Font(name="Segoe UI", size=10, bold=True, color="1E293B")

        thin_side = Side(border_style="thin", color="CBD5E1")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        ws.merge_cells("A1:F1")
        ws["A1"] = "AUDIT TRAIL LOG ? SISTEM MONITORING DOKUMEN PROYEK"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(vertical="center")

        ws.merge_cells("A2:F2")
        ws["A2"] = "PT PAKARTI RIKEN INDONESIA ? MACHINING DIVISION"
        ws["A2"].font = subtitle_font
        ws["A2"].alignment = Alignment(vertical="center")

        filter_info = f"Exported: {datetime.now().strftime('%d-%b-%Y %H:%M:%S')} | Total Records: {len(logs)}"
        if start_date or end_date:
            filter_info += f" | Filter Range: {start_date or 'Awal'} s/d {end_date or 'Sekarang'}"
        if search:
            filter_info += f" | Keyword: '{search}'"

        ws.merge_cells("A3:F3")
        ws["A3"] = filter_info
        ws["A3"].font = Font(name="Segoe UI", size=9, color="475569")
        ws["A3"].alignment = Alignment(vertical="center")

        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 18
        ws.row_dimensions[3].height = 18
        ws.row_dimensions[4].height = 10
        ws.row_dimensions[5].height = 26

        headers = ["No", "Timestamp", "User ID / Role", "Nama User", "Action", "Details"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = navy_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for idx, log in enumerate(logs, 1):
            row_idx = idx + 5
            ws.row_dimensions[row_idx].height = 20

            ts_str = log.timestamp
            try:
                dt = datetime.fromisoformat(log.timestamp)
                ts_str = dt.strftime("%Y-%m-%d %H:%M:%S")
            except Exception:
                pass

            row_data = [
                idx,
                ts_str,
                log.user_id or "-",
                log.user_name or "Guest",
                log.action or "-",
                log.details or "-"
            ]

            fill_to_use = accent_fill if idx % 2 == 0 else PatternFill(fill_type=None)

            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = data_font
                cell.border = thin_border
                if fill_to_use.fill_type:
                    cell.fill = fill_to_use

                if col_idx == 1:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx in (2, 3):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 4:
                    cell.font = bold_data_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif col_idx == 5:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        col_widths = {1: 8, 2: 22, 3: 18, 4: 26, 5: 24, 6: 70}
        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"Audit_Log_Monitoring_System_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename
        )

    except Exception as e:
        return jsonify({'error': str(e)}), 500
