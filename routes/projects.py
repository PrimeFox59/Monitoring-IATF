from flask import Blueprint, jsonify, request, send_file, current_app
from models import db, ProjectNew, ProjectDocument, DocumentType, User
from routes.audit import log_activity
from routes.documents import get_absolute_path
from sqlalchemy.orm import joinedload
from datetime import datetime
try:
    import pandas as pd
except ImportError:
    pd = None
import openpyxl
import io
import os
import zipfile

projects_bp = Blueprint('projects', __name__)

@projects_bp.route('/api/projects', methods=['GET'])
def get_projects():
    """Get all projects with basic info - optimized with eager loading"""
    try:
        page = request.args.get('page', 1, type=int)
        per_page = request.args.get('per_page', 100, type=int)
        search = request.args.get('search', '').strip()
        filter_project = request.args.get('filter_project')
        filter_customer = request.args.get('filter_customer')
        filter_status = request.args.get('filter_status')
        filter_department = request.args.get('filter_department')
        filter_department = request.args.get('filter_department')

        projects_query = ProjectNew.query.options(joinedload(ProjectNew.documents)).order_by(
            ProjectNew.updated_at.desc(), ProjectNew.id.desc()
        )

        if search:
            search_query = f"%{search}%"
            projects_query = projects_query.filter(
                (ProjectNew.item.ilike(search_query)) |
                (ProjectNew.part_no.ilike(search_query)) |
                (ProjectNew.i_code.ilike(search_query)) |
                (ProjectNew.project.ilike(search_query)) |
                (ProjectNew.customer.ilike(search_query)) |
                (ProjectNew.pic.ilike(search_query))
            )

        if filter_project:
            projects_query = projects_query.filter(ProjectNew.project == filter_project)
        if filter_customer:
            projects_query = projects_query.filter(ProjectNew.customer == filter_customer)
        if filter_status:
            projects_query = projects_query.filter(ProjectNew.status == filter_status)

        paginated = projects_query.paginate(page=page, per_page=per_page, error_out=False)

        all_dtypes = DocumentType.query.order_by(DocumentType.id).all()
        dtype_obj_map = {dt.id: dt for dt in all_dtypes}

        if filter_department:
            scoped_dtypes = [dt for dt in all_dtypes if (dt.department or 'Prod MC') == filter_department]
            scoped_ids = {dt.id for dt in scoped_dtypes}
        else:
            scoped_dtypes = all_dtypes
            scoped_ids = None

        result = []
        for project in paginated.items:
            all_p_docs = project.documents
            if scoped_ids is not None:
                docs = [d for d in all_p_docs if d.doc_type_id in scoped_ids or (d.department or 'Prod MC') == filter_department]
            else:
                docs = all_p_docs

            total_docs = len(scoped_dtypes)
            uploaded_docs = 0
            doc_matrix = {}
            missing_docs = []

            # Map doc_type_id to uploaded status
            d_up_map = {}
            for d in docs:
                if d.status in ('Uploaded', 'Completed'):
                    d_up_map[d.doc_type_id] = True

            for dt in scoped_dtypes:
                is_up = d_up_map.get(dt.id, False)
                doc_matrix[dt.name] = is_up
                if is_up:
                    uploaded_docs += 1
                else:
                    missing_docs.append(dt.name)

            progress = int((uploaded_docs / total_docs) * 100) if total_docs > 0 else 0

            # Collect distinct departments across all documents in this project
            proj_depts = sorted(list({(dtype_obj_map[d.doc_type_id].department if d.doc_type_id in dtype_obj_map else d.department) or 'Prod MC' for d in all_p_docs} or {'Prod MC'}))

            p_dict = {
                'id': project.id,
                'item': project.item or '',
                'part_no': project.part_no or '',
                'project_name': project.project or '',
                'customer': project.customer or '',
                'status': project.status or 'On Progress',
                'pic': project.pic or '',
                'i_code': project.i_code or '',
                'project_start_date': project.project_start_date or '',
                'project_end_date': project.project_end_date or '',
                'department': filter_department or ', '.join(proj_depts),
                'departments': proj_depts,
                'total_docs': total_docs,
                'uploaded_docs': uploaded_docs,
                'progress': progress,
                'doc_matrix': doc_matrix,
                'missing_docs': missing_docs
            }
            result.append(p_dict)

        return jsonify({
            'items': result,
            'total': paginated.total,
            'page': page,
            'per_page': per_page,
            'pages': paginated.pages
        }), 200

    except Exception as e:
        return jsonify({'error': str(e)}), 500

@projects_bp.route('/api/projects/export-matrix', methods=['GET'])
def export_matrix():
    """Export executive IATF 16949 Document Compliance Matrix to styled Excel workbook"""
    try:
        search = request.args.get('search', '').strip()
        filter_project = request.args.get('filter_project')
        filter_customer = request.args.get('filter_customer')
        filter_status = request.args.get('filter_status')
        filter_department = request.args.get('filter_department')

        projects_query = ProjectNew.query.options(joinedload(ProjectNew.documents)).order_by(
            ProjectNew.updated_at.desc(), ProjectNew.id.desc()
        )

        if search:
            search_query = f"%{search}%"
            projects_query = projects_query.filter(
                (ProjectNew.item.ilike(search_query)) |
                (ProjectNew.part_no.ilike(search_query)) |
                (ProjectNew.i_code.ilike(search_query)) |
                (ProjectNew.project.ilike(search_query)) |
                (ProjectNew.customer.ilike(search_query)) |
                (ProjectNew.pic.ilike(search_query))
            )

        if filter_project:
            projects_query = projects_query.filter(ProjectNew.project == filter_project)
        if filter_customer:
            projects_query = projects_query.filter(ProjectNew.customer == filter_customer)
        if filter_status:
            projects_query = projects_query.filter(ProjectNew.status == filter_status)

        projects = projects_query.all()
        all_dtypes = DocumentType.query.order_by(DocumentType.id).all()

        if filter_department:
            scoped_dtypes = [dt for dt in all_dtypes if (dt.department or 'Prod MC') == filter_department]
        else:
            scoped_dtypes = all_dtypes

        dtype_map = {dt.id: dt.name for dt in scoped_dtypes}
        doc_names = [dt.name for dt in scoped_dtypes]

        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Matriks Dokumen Proyek"

        # Executive color styling (Deep Navy, Emerald & Rose)
        navy_dark = "18345E"
        navy_header = "283C5A"
        green_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
        green_font = Font(name="Segoe UI", size=9, bold=True, color="065F46")
        red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        red_font = Font(name="Segoe UI", size=9, bold=True, color="991B1B")

        title_font = Font(name="Segoe UI", size=13, bold=True, color=navy_dark)
        subtitle_font = Font(name="Segoe UI", size=9, italic=True, color="475569")
        th_font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
        th_fill = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
        th_doc_fill = PatternFill(start_color=navy_header, end_color=navy_header, fill_type="solid")
        cell_font = Font(name="Segoe UI", size=9, color="1E293B")

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        dept_label = filter_department.upper() if filter_department else "MACHINING DIVISION"
        ws['A1'] = f"PT PAKARTI RIKEN INDONESIA ??? {dept_label}"
        ws['A1'].font = Font(name="Segoe UI", size=10, bold=True, color="64748B")
        ws['A2'] = "MATRIKS KELENGKAPAN DOKUMEN PROYEK"
        ws['A2'].font = title_font
        ws['A3'] = f"Waktu Export: {datetime.now().strftime('%d-%m-%Y %H:%M:%S')} WIB | Total Proyek Terfilter: {len(projects)}"
        ws['A3'].font = subtitle_font

        headers = ["No", "I-Code", "Item", "Part No", "Project", "Customer", "PIC", "Status", "Progress", "Ada", "Total"]
        all_headers = headers + doc_names

        for col_num, h_text in enumerate(all_headers, 1):
            cell = ws.cell(row=5, column=col_num, value=h_text)
            cell.font = th_font
            cell.fill = th_fill if col_num <= len(headers) else th_doc_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        for r_idx, proj in enumerate(projects, 6):
            docs = proj.documents
            total_d = len(docs)
            uploaded_d = 0
            d_map = {}
            for d in docs:
                name = dtype_map.get(d.doc_type_id)
                if name:
                    is_up = (d.status == 'Uploaded')
                    d_map[name] = is_up
                    if is_up:
                        uploaded_d += 1
            prog = int((uploaded_d / total_d) * 100) if total_d > 0 else 0

            row_data = [
                r_idx - 5,
                proj.i_code or "-",
                proj.item or "-",
                proj.part_no or "-",
                proj.project or "-",
                proj.customer or "-",
                proj.pic or "-",
                proj.status or "-",
                f"{prog}%",
                uploaded_d,
                total_d
            ]

            for c_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = cell_font
                cell.border = thin_border
                cell.alignment = Alignment(
                    horizontal="center" if c_idx in [1, 2, 8, 9, 10, 11] else "left",
                    vertical="center"
                )

            for d_idx, doc_name in enumerate(doc_names, len(headers) + 1):
                has_doc = d_map.get(doc_name, False)
                cell = ws.cell(row=r_idx, column=d_idx)
                if has_doc:
                    cell.value = "ADA"
                    cell.fill = green_fill
                    cell.font = green_font
                else:
                    cell.value = "KOSONG"
                    cell.fill = red_fill
                    cell.font = red_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

        ws.row_dimensions[5].height = 28
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"Matriks_Monitoring_System_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename
        )

    except Exception as e:
        return jsonify({'error': str(e)}), 500


@projects_bp.route('/api/projects', methods=['POST'])
def create_project():
    """Create a new project"""
    try:
        data = request.get_json(silent=True) or request.json or {}
        if data.get('i_code'):
            if ProjectNew.query.filter_by(i_code=data.get('i_code')).first():
                return jsonify({'error': 'Project with this I-Code already exists'}), 400

        new_project = ProjectNew(
            i_code=data.get('i_code'),
            project=data.get('project_name') or data.get('project'),
            item=data.get('item'),
            part_no=data.get('part_no'),
            customer=data.get('customer'),
            department=data.get('department') or 'Prod MC',
            pic=data.get('pic'),
            status=data.get('status', 'On Progress'),
            project_start_date=data.get('project_start_date'),
            project_end_date=data.get('project_end_date'),
            created_at=datetime.now().isoformat(),
            updated_at=datetime.now().isoformat()
        )
        db.session.add(new_project)
        db.session.flush()

        all_types = DocumentType.query.all()
        for dtype in all_types:
            new_doc = ProjectDocument(
                project_id=new_project.id,
                doc_type_id=dtype.id,
                status='Empty',
                delegated_to_id=dtype.default_delegate_id
            )
            db.session.add(new_doc)

        db.session.commit()

        admin_name = request.headers.get('X-User-Name', 'Admin')
        log_activity(
            request.headers.get('X-User-Role', 'System'),
            admin_name,
            'CREATE_PROJECT',
            f"Created Project: {new_project.project} | Item: {new_project.item or '-'} | Part No: {new_project.part_no or '-'} | Customer: {new_project.customer or '-'}"
        )

        return jsonify(new_project.to_dict()), 201

    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500

@projects_bp.route('/api/projects/<int:project_id>', methods=['GET'])
def get_project_detail(project_id):
    """Get project details with full documents and revisions"""
    try:
        project = ProjectNew.query.get_or_404(project_id)
        return jsonify(project.to_dict(include_documents=True)), 200

    except Exception as e:
        return jsonify({'error': str(e)}), 500

@projects_bp.route('/api/projects/<int:project_id>', methods=['PUT'])
def update_project(project_id):
    """Update an existing project"""
    try:
        project = ProjectNew.query.get_or_404(project_id)
        data = request.get_json(silent=True) or request.json or {}

        if 'i_code' in data:
            if data['i_code'] and data['i_code'] != project.i_code:
                if ProjectNew.query.filter_by(i_code=data['i_code']).first():
                    return jsonify({'error': 'Project with this I-Code already exists'}), 400
                project.i_code = data['i_code']

        if 'project_name' in data:
            project.project = data['project_name']
        elif 'project' in data:
            project.project = data['project']

        if 'item' in data:
            project.item = data['item']
        if 'part_no' in data:
            project.part_no = data['part_no']
        if 'customer' in data:
            project.customer = data['customer']
        if 'department' in data:
            project.department = data['department'] or 'Prod MC'
        if 'pic' in data:
            project.pic = data['pic']
        if 'status' in data:
            project.status = data['status']
        if 'project_start_date' in data:
            project.project_start_date = data['project_start_date']
        if 'project_end_date' in data:
            project.project_end_date = data['project_end_date']

        project.updated_at = datetime.now().isoformat()
        db.session.commit()

        admin_name = request.headers.get('X-User-Name', 'Admin')
        log_activity(
            request.headers.get('X-User-Role', 'System'),
            admin_name,
            'UPDATE_PROJECT',
            f"Updated Project: {project.project} | Item: {project.item or '-'} | Part No: {project.part_no or '-'} | Customer: {project.customer or '-'}"
        )

        return jsonify(project.to_dict()), 200

    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500

@projects_bp.route('/api/projects/<int:project_id>', methods=['DELETE'])
def delete_project(project_id):
    """Delete a project and its documents"""
    try:
        project = ProjectNew.query.get_or_404(project_id)
        project_name = project.project
        project_info = f"Project: {project_name} | Item: {project.item or '-'} | Part No: {project.part_no or '-'} | Customer: {project.customer or '-'}"

        for doc in project.documents:
            if doc.file_path and os.path.exists(doc.file_path):
                try:
                    os.remove(doc.file_path)
                except Exception:
                    pass
            for rev in doc.revisions:
                if rev.file_path and os.path.exists(rev.file_path):
                    try:
                        os.remove(rev.file_path)
                    except Exception:
                        pass

        db.session.delete(project)
        db.session.commit()

        admin_name = request.headers.get('X-User-Name', 'Admin')
        log_activity(
            request.headers.get('X-User-Role', 'System'),
            admin_name,
            'DELETE_PROJECT',
            f"Deleted {project_info}"
        )

        return jsonify({'message': 'Project deleted successfully'}), 200

    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500

@projects_bp.route('/api/stats', methods=['GET'])
@projects_bp.route('/api/projects/stats', methods=['GET'])
def get_stats():
    """Get dashboard stats matching frontend expectations"""
    try:
        total_projects = ProjectNew.query.count()
        done = ProjectNew.query.filter((ProjectNew.status == 'Done') | (ProjectNew.status == 'Completed')).count()
        on_progress = ProjectNew.query.filter((ProjectNew.status == 'On Progress') | (ProjectNew.status == 'In Progress')).count()
        hold = ProjectNew.query.filter_by(status='Hold').count()
        canceled = ProjectNew.query.filter((ProjectNew.status == 'Canceled') | (ProjectNew.status == 'Cancelled')).count()

        total_docs = ProjectDocument.query.count()
        uploaded_docs = ProjectDocument.query.filter_by(status='Uploaded').count()
        overall_progress = int((uploaded_docs / total_docs) * 100) if total_docs > 0 else 0

        return jsonify({
            'total': total_projects,
            'total_projects': total_projects,
            'on_progress': on_progress,
            'done': done,
            'completed': done,
            'hold': hold,
            'canceled': canceled,
            'total_documents': total_docs,
            'uploaded_documents': uploaded_docs,
            'overall_progress': overall_progress
        }), 200

    except Exception as e:
        return jsonify({'error': str(e)}), 500

@projects_bp.route('/api/projects/metadata', methods=['GET'])
def get_projects_metadata():
    """Get metadata for filtering: projects, customers, pics, and statuses"""
    try:
        projects = ProjectNew.query.all()
        customers = sorted(list({p.customer for p in projects if p.customer}))
        project_names = sorted(list({p.project for p in projects if p.project}))
        pics = sorted(list({p.pic for p in projects if p.pic}))
        statuses = sorted(list({p.status for p in projects if p.status}))
        all_dt_depts = {dt.department for dt in DocumentType.query.all() if dt.department}
        departments = sorted(list(all_dt_depts or {'Prod MC'}))
        if 'Prod MC' not in departments:
            departments.insert(0, 'Prod MC')

        return jsonify({
            'projects': project_names,
            'project_names': project_names,
            'customers': customers,
            'pics': pics,
            'statuses': statuses,
            'departments': departments
        }), 200

    except Exception as e:
        return jsonify({'error': str(e)}), 500

@projects_bp.route('/api/projects/import-excel', methods=['POST'])
def import_projects_excel():
    """Import projects from uploaded Excel file"""
    if 'file' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400

    file = request.files['file']
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({'error': 'Invalid file format. Please upload an Excel file.'}), 400

    try:
        imported = 0
        skipped = 0
        all_types = DocumentType.query.all()

        rows_data = []
        if pd is not None:
            df = pd.read_excel(file)
            for _, row in df.iterrows():
                rows_data.append({
                    'i_code': str(row.get('I-Code', '')).strip() if pd.notna(row.get('I-Code')) else None,
                    'project_name': str(row.get('Project Name', row.get('Project', ''))).strip() if pd.notna(row.get('Project Name', row.get('Project'))) else None,
                    'item': str(row.get('Item', '')).strip() if pd.notna(row.get('Item')) else None,
                    'part_no': str(row.get('Part No', '')).strip() if pd.notna(row.get('Part No')) else None,
                    'customer': str(row.get('Customer', '')).strip() if pd.notna(row.get('Customer')) else None,
                    'pic': str(row.get('PIC', '')).strip() if pd.notna(row.get('PIC')) else None,
                    'status': str(row.get('Status', 'On Progress')).strip() if pd.notna(row.get('Status')) else 'On Progress'
                })
        else:
            wb = openpyxl.load_workbook(file)
            ws = wb.active
            header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
            headers = [str(c).strip().lower() if c is not None else '' for c in header_row]
            header_map = {h: idx for idx, h in enumerate(headers) if h}

            for row in ws.iter_rows(min_row=2, values_only=True):
                if not any(row):
                    continue
                def get_c(aliases):
                    for a in aliases:
                        idx = header_map.get(a.lower())
                        if idx is not None and idx < len(row):
                            v = row[idx]
                            if v is not None and str(v).strip():
                                return str(v).strip()
                    return None

                rows_data.append({
                    'i_code': get_c(['i-code', 'icode', 'i_code']),
                    'project_name': get_c(['project name', 'project', 'project_name']),
                    'item': get_c(['item']),
                    'part_no': get_c(['part no', 'part_no']),
                    'customer': get_c(['customer']),
                    'pic': get_c(['pic']),
                    'status': get_c(['status']) or 'On Progress'
                })

        for row_dict in rows_data:
            i_code = row_dict['i_code']
            project_name = row_dict['project_name']
            item = row_dict['item']
            part_no = row_dict['part_no']
            customer = row_dict['customer']
            pic = row_dict['pic']
            status = row_dict['status']

            if not project_name and not item:
                skipped += 1
                continue

            if i_code and ProjectNew.query.filter_by(i_code=i_code).first():
                skipped += 1
                continue

            new_project = ProjectNew(
                i_code=i_code,
                project=project_name,
                item=item,
                part_no=part_no,
                customer=customer,
                pic=pic,
                status=status,
                created_at=datetime.now().isoformat(),
                updated_at=datetime.now().isoformat()
            )
            db.session.add(new_project)
            db.session.flush()

            for dtype in all_types:
                new_doc = ProjectDocument(
                    project_id=new_project.id,
                    doc_type_id=dtype.id,
                    status='Empty',
                    delegated_to_id=dtype.default_delegate_id
                )
                db.session.add(new_doc)

            imported += 1

        db.session.commit()

        admin_name = request.headers.get('X-User-Name', 'Admin')
        log_activity(
            request.headers.get('X-User-Role', 'System'),
            admin_name,
            'IMPORT_EXCEL',
            f'Imported {imported} projects, skipped {skipped}'
        )

        return jsonify({
            'message': f'Successfully imported {imported} projects ({skipped} skipped)',
            'imported': imported,
            'skipped': skipped
        }), 200

    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 500

@projects_bp.route('/api/projects/<int:project_id>/download-all', methods=['GET'])
def download_project_documents(project_id):
    """Download all uploaded documents of a project in a zip file"""
    try:
        project = ProjectNew.query.get_or_404(project_id)
        memory_file = io.BytesIO()

        with zipfile.ZipFile(memory_file, 'w', zipfile.ZIP_DEFLATED) as zf:
            for doc in project.documents:
                doc_type_name = doc.doc_type.name if doc.doc_type else 'Document'
                if doc.doc_type and doc.doc_type.is_multiple and doc.revisions:
                    for rev in doc.revisions:
                        abs_p = get_absolute_path(rev.file_path)
                        if abs_p and os.path.exists(abs_p):
                            safe_name = f"{doc_type_name}_Rev{rev.revision_number}_{os.path.basename(abs_p)}"
                            zf.write(abs_p, arcname=safe_name)
                elif doc.file_path:
                    abs_p = get_absolute_path(doc.file_path)
                    if abs_p and os.path.exists(abs_p):
                        safe_name = f"{doc_type_name}_{os.path.basename(abs_p)}"
                        zf.write(abs_p, arcname=safe_name)

        memory_file.seek(0)
        zip_filename = f"{project.project or 'Project'}_{project.id}_documents.zip"

        return send_file(
            memory_file,
            mimetype='application/zip',
            as_attachment=True,
            download_name=zip_filename
        )

    except Exception as e:
        return jsonify({'error': str(e)}), 500


